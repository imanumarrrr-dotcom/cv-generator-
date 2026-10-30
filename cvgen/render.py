"""One CV, four templates, five export formats (HTML, PDF, DOCX, Markdown, JSON)."""
import html as H, io, json
from .models import clean
T = {  # name: dark, bg, fg, muted, a, b, head font, body font, layout, pdf font, pdf accent, pdf accent2
 "Human Next": dict(dark=True, bg="#0A0A0A", fg="#EEF3F8", muted="#8FA3B8", a="#2FB0E3", b="#FF2D95", hf="'Michroma','Segoe UI',sans-serif", bf="'Manrope','Segoe UI',sans-serif", layout="one", pf="Helvetica", pa="#007BA7", pb="#FF2D95"),
 "Executive":  dict(dark=False, bg="#FFFFFF", fg="#1B2335", muted="#5B6577", a="#1B2A49", b="#8A6D3B", hf="Georgia,'Times New Roman',serif", bf="Georgia,'Times New Roman',serif", layout="one", pf="Times-Roman", pa="#1B2A49", pb="#8A6D3B"),
 "Modern":     dict(dark=False, bg="#FFFFFF", fg="#10202B", muted="#5A6B78", a="#007BA7", b="#FF2D95", hf="'Manrope','Segoe UI',sans-serif", bf="'Manrope','Segoe UI',sans-serif", layout="two", pf="Helvetica", pa="#007BA7", pb="#FF2D95"),
 "Minimal":    dict(dark=False, bg="#FFFFFF", fg="#111111", muted="#666666", a="#111111", b="#111111", hf="'Manrope','Segoe UI',sans-serif", bf="'Manrope','Segoe UI',sans-serif", layout="one", pf="Helvetica", pa="#111111", pb="#999999")}
e = lambda s: H.escape(s or "")
def _dates(x): return " – ".join(p for p in (x["start"], x["end"]) if p)
def _contact(c): return [x for x in (c["email"], c["phone"], c["location"]) + tuple(f'{l["label"]}: {l["url"]}' if l["label"] else l["url"] for l in c["links"]) if x]

def to_html(cv, tpl="Human Next"):
    c, t = clean(cv), T[tpl]; sec = lambda h, b: f'<section><h2>{h}</h2>{b}</section>' if b.strip() else ""
    exp = "".join(f'<div class="item"><div class="row"><b>{e(x["role"])}</b><span class="d">{e(_dates(x))}</span></div><div class="m">{e(", ".join(p for p in (x["company"], x["location"]) if p))}</div><ul>{"".join(f"<li>{e(b)}</li>" for b in x["bullets"])}</ul></div>' for x in c["experience"])
    edu = "".join(f'<div class="item"><div class="row"><b>{e(x["degree"])}</b><span class="d">{e(_dates(x))}</span></div><div class="m">{e(", ".join(p for p in (x["school"], x["location"]) if p))}</div><p>{e(x["details"])}</p></div>' for x in c["education"])
    prj = "".join(f'<div class="item"><div class="row"><b>{e(x["name"])}</b><span class="d">{e(x["link"])}</span></div><p>{e(x["description"])} <i>{e(x["tech"])}</i></p></div>' for x in c["projects"])
    skl = "".join(f'<span class="chip">{e(s)}</span>' for s in c["skills"]); lst = lambda xs: "<ul>" + "".join(f"<li>{e(x)}</li>" for x in xs) + "</ul>" if xs else ""
    head = f'<header><h1>{e(c["name"]) or "Your Name"}</h1><div class="hl">{e(c["headline"])}</div><div class="ct">{" · ".join(e(x) for x in _contact(c))}</div></header>'
    main = sec("Summary", f'<p>{e(c["summary"])}</p>' if c["summary"] else "") + sec("Experience", exp) + sec("Projects", prj) + sec("Education", edu)
    side = sec("Skills", f'<div class="chips">{skl}</div>') + sec("Certifications", lst(c["certifications"])) + sec("Languages", lst(c["languages"]))
    body = f'<div class="grid"><aside>{side}</aside><main>{main}</main></div>' if t["layout"] == "two" else f"<main>{main}</main>{side}"
    glow = "text-shadow:0 0 18px color-mix(in srgb,var(--b) 55%,transparent);" if t["dark"] else ""
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>{e(c["name"])} - CV</title><link href="https://fonts.googleapis.com/css2?family=Michroma&family=Manrope:wght@400;600;800&display=swap" rel="stylesheet">
<style>:root{{--a:{t["a"]};--b:{t["b"]};--bg:{t["bg"]};--fg:{t["fg"]};--m:{t["muted"]}}}*{{box-sizing:border-box}}body{{margin:0;background:{"#050505" if t["dark"] else "#EEF1F5"};font:14px/1.55 {t["bf"]};color:var(--fg)}}
.page{{max-width:820px;margin:24px auto;background:var(--bg);padding:44px 48px;{"box-shadow:0 0 60px rgba(47,176,227,.12);border:1px solid #1b2a33;" if t["dark"] else "box-shadow:0 6px 30px rgba(0,0,0,.12);"}}}
h1{{font:700 28px {t["hf"]};letter-spacing:.04em;margin:0;{glow}background:{"linear-gradient(90deg,var(--a),var(--b))" if t["dark"] else "none"};-webkit-background-clip:text;background-clip:text;color:{"transparent" if t["dark"] else "var(--a)"}}}
.hl{{color:var(--b);font-weight:600;margin-top:4px}}.ct{{color:var(--m);font-size:12px;margin:6px 0 18px}}header{{border-bottom:2px solid var(--a);margin-bottom:6px}}
h2{{font:700 11px {t["hf"]};letter-spacing:.28em;text-transform:uppercase;color:var(--a);margin:20px 0 8px;display:flex;align-items:center;gap:10px}}h2:after{{content:"";flex:1;height:1px;background:linear-gradient(90deg,var(--b),transparent)}}
.row{{display:flex;justify-content:space-between;gap:12px}}.d{{color:var(--m);font-size:12px;white-space:nowrap}}.m{{color:var(--m);font-size:13px}}.item{{margin-bottom:12px}}ul{{margin:6px 0 0;padding-left:18px}}li{{margin:2px 0}}p{{margin:4px 0}}
.chip{{display:inline-block;margin:0 6px 6px 0;padding:2px 10px;border:1px solid var(--a);border-radius:99px;font-size:12px}}.grid{{display:grid;grid-template-columns:1fr 2.2fr;gap:28px}}aside{{border-right:1px solid var(--b);padding-right:18px}}
@media print{{body{{background:#fff}}.page{{box-shadow:none;margin:0;border:0}}}}@media(max-width:700px){{.grid{{grid-template-columns:1fr}}.page{{padding:24px}}}}</style></head><body><div class="page">{head}{body}</div></body></html>"""

def to_md(cv):
    c = clean(cv); o = [f"# {c['name']}", f"**{c['headline']}**", " | ".join(_contact(c)), "", "## Summary", c["summary"], "", "## Experience"]
    for x in c["experience"]: o += [f"### {x['role']} — {x['company']}", f"*{_dates(x)} · {x['location']}*"] + [f"- {b}" for b in x["bullets"]] + [""]
    o += ["## Education"] + [f"- **{x['degree']}**, {x['school']} ({_dates(x)}) {x['details']}" for x in c["education"]]
    o += ["", "## Skills", ", ".join(c["skills"])]
    if c["projects"]: o += ["", "## Projects"] + [f"- **{x['name']}** — {x['description']} ({x['tech']})" for x in c["projects"]]
    if c["certifications"]: o += ["", "## Certifications"] + [f"- {x}" for x in c["certifications"]]
    if c["languages"]: o += ["", "## Languages", ", ".join(c["languages"])]
    return "\n".join(o)

def to_docx(cv):
    import docx
    from docx.shared import Pt, RGBColor
    c = clean(cv); d = docx.Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)
    def H2(t):
        p = d.add_paragraph(); r = p.add_run(t.upper()); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x00, 0x7B, 0xA7)
    p = d.add_paragraph(); r = p.add_run(c["name"] or "Your Name"); r.bold = True; r.font.size = Pt(22)
    if c["headline"]: d.add_paragraph().add_run(c["headline"]).bold = True
    if _contact(c): d.add_paragraph(" | ".join(_contact(c)))
    if c["summary"]: H2("Summary"); d.add_paragraph(c["summary"])
    if c["experience"]: H2("Experience")
    for x in c["experience"]:
        p = d.add_paragraph(); r = p.add_run(f"{x['role']} — {x['company']}"); r.bold = True; p.add_run(f"   {_dates(x)}")
        for b in x["bullets"]: d.add_paragraph(b, style="List Bullet")
    if c["education"]: H2("Education")
    for x in c["education"]: d.add_paragraph(f"{x['degree']}, {x['school']} ({_dates(x)}) {x['details']}".strip())
    if c["projects"]: H2("Projects")
    for x in c["projects"]: d.add_paragraph(f"{x['name']} — {x['description']} ({x['tech']})")
    if c["skills"]: H2("Skills"); d.add_paragraph(", ".join(c["skills"]))
    for t, k in (("Certifications", "certifications"), ("Languages", "languages")):
        if c[k]: H2(t); d.add_paragraph(", ".join(c[k]))
    b = io.BytesIO(); d.save(b); return b.getvalue()

def to_pdf(cv, tpl="Human Next"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    t, c = T[tpl], clean(cv); a, b = colors.HexColor(t["pa"]), colors.HexColor(t["pb"]); f = t["pf"]; fb = f + ("-Bold" if f == "Helvetica" else "")
    fb = "Times-Bold" if f == "Times-Roman" else fb; esc = lambda s: H.escape(s or "").replace("—", "-")
    S = lambda n, **k: ParagraphStyle(n, fontName=k.pop("fontName", f), fontSize=k.pop("fontSize", 9.5), leading=k.pop("leading", 13), **k)
    nm, hl, ct = S("n", fontName=fb, fontSize=22, leading=26, textColor=a), S("h", fontName=fb, fontSize=10.5, textColor=b), S("c", fontSize=8.5, textColor=colors.grey, spaceAfter=4)
    h2, it, mt, bd = S("h2", fontName=fb, fontSize=9, textColor=a, spaceBefore=9, spaceAfter=2), S("i", fontName=fb, spaceBefore=4), S("m", fontSize=8.5, textColor=colors.grey), S("b", leftIndent=10, bulletIndent=0)
    out = [Paragraph(esc(c["name"]), nm), Paragraph(esc(c["headline"]), hl), Paragraph(" · ".join(esc(x) for x in _contact(c)), ct), HRFlowable(width="100%", thickness=1.4, color=a)]
    def sec(title):
        out.extend([Paragraph(title.upper(), h2), HRFlowable(width="100%", thickness=.4, color=b, spaceAfter=3)])
    if c["summary"]: sec("Summary"); out.append(Paragraph(esc(c["summary"]), bd.clone("s", leftIndent=0)))
    if c["experience"]: sec("Experience")
    for x in c["experience"]:
        out += [Paragraph(f"{esc(x['role'])} - {esc(x['company'])}", it), Paragraph(esc(", ".join(p for p in (_dates(x), x["location"]) if p)), mt)] + [Paragraph(esc(bl), bd, bulletText="•") for bl in x["bullets"]]
    if c["projects"]: sec("Projects"); out += [Paragraph(f"<b>{esc(x['name'])}</b> - {esc(x['description'])} <i>{esc(x['tech'])}</i>", bd.clone("p", leftIndent=0)) for x in c["projects"]]
    if c["education"]: sec("Education"); out += [Paragraph(f"<b>{esc(x['degree'])}</b>, {esc(x['school'])} ({esc(_dates(x))}) {esc(x['details'])}", bd.clone("e", leftIndent=0)) for x in c["education"]]
    if c["skills"]: sec("Skills"); out.append(Paragraph(esc(", ".join(c["skills"])), bd.clone("k", leftIndent=0)))
    for ttl, k in (("Certifications", "certifications"), ("Languages", "languages")):
        if c[k]: sec(ttl); out.append(Paragraph(esc(", ".join(c[k])), bd.clone(k, leftIndent=0)))
    buf = io.BytesIO(); SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=f"{c['name']} CV").build(out); return buf.getvalue()

def to_json(cv): return json.dumps(clean(cv), indent=2, ensure_ascii=False)
